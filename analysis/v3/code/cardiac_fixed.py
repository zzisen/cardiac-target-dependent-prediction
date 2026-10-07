# Project-authored code: MIT; scientific functions preserved verbatim.
from __future__ import annotations

import csv

import hashlib

import json

import math

import platform

import re

import sys

from datetime import datetime, timezone

from pathlib import Path

import numpy as np

import pandas as pd

from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent

ROOT = Path(__file__).resolve().parents[1] / 'inputs' / 'project'

RLC_ROOT = Path(__file__).resolve().parents[1] / 'inputs' / 'project' / 'P2_RLC1_DUAL_Q1_UPLIFT_2026-10-03'

V2_ROOT = Path(__file__).resolve().parents[3]

RLC_FAMILIES = ["PeakTension", "TTP", "RT50"]

RLC_CANDIDATES = [f"1uM_{family}" for family in RLC_FAMILIES]

RLC_TARGETS = [f"{dose}_{family}" for dose in ("3uM", "10uM") for family in RLC_FAMILIES]

RLC_SHEETS = {
    "PeakTension": "Fig6B_TPeak",
    "TTP": "Fig6D_TTP",
    "RT50": "Fig6E_RT50",
}

AW_TARGETS = [0.05, 0.10, 0.25, 0.50, 1.00, 2.50]

BASE_ATP = 5.0

ALPHA = 1.0

def _sample_sd(values: np.ndarray, label: str) -> float:
    sd = float(np.std(np.asarray(values, dtype=float), ddof=1))
    if not math.isfinite(sd) or sd <= 0.0:
        raise RuntimeError(f"Nonpositive/nonfinite fold-local sample SD: {label}")
    return sd

def fit_rlc_standardized_ols(
    x: np.ndarray,
    y: np.ndarray,
    train: np.ndarray,
    test: np.ndarray,
    label: str,
) -> tuple[np.ndarray, float, float, float, float]:
    x_mu = float(np.mean(x[train]))
    x_sd = _sample_sd(x[train], label + " candidate")
    y_mu = float(np.mean(y[train]))
    y_sd = _sample_sd(y[train], label + " target")
    xz_train = (x[train] - x_mu) / x_sd
    yz_train = (y[train] - y_mu) / y_sd
    design = np.column_stack((np.ones(len(train), dtype=float), xz_train))
    beta, _, rank, _ = np.linalg.lstsq(design, yz_train, rcond=None)
    if rank != 2 or not np.isfinite(beta).all():
        raise RuntimeError(f"RLC affine OLS rank/finite failure: {label}")
    xz_test = (x[test] - x_mu) / x_sd
    pred_z = beta[0] + beta[1] * xz_test
    pred_native = y_mu + y_sd * pred_z
    return pred_native, y_mu, y_sd, x_mu, x_sd

def load_rlc_support() -> tuple[dict[str, np.ndarray], dict[str, np.ndarray]]:
    source = RLC_ROOT / "RLC1_SOURCE_DATA" / "RLC-1_Fig6.xlsx"
    wb = load_workbook(source, read_only=True, data_only=True)
    candidate_values: dict[str, np.ndarray] = {}
    target_values: dict[str, np.ndarray] = {}
    for family, sheet_name in RLC_SHEETS.items():
        rows = list(wb[sheet_name].iter_rows(values_only=True))
        header = tuple(str(v).strip() for v in rows[0])
        body = rows[1:]
        if header != ("ND", "1uM", "3uM", "10uM") or len(body) != 7 or any(len(r) != 4 for r in body):
            raise RuntimeError(f"RLC frozen workbook shape/header mismatch: {sheet_name}")
        values = np.asarray(body, dtype=float)
        if values.shape != (7, 4) or not np.isfinite(values).all():
            raise RuntimeError(f"RLC incomplete/nonfinite rat rows: {sheet_name}")
        candidate_values[f"1uM_{family}"] = values[:, 1]
        target_values[f"3uM_{family}"] = values[:, 2]
        target_values[f"10uM_{family}"] = values[:, 3]
        relevant = values[:, [1, 2, 3]]
        for outer in range(7):
            six = np.delete(relevant, outer, axis=0)
            if not np.all(np.std(six, axis=0, ddof=1) > 0.0):
                raise RuntimeError(f"RLC zero outer fit SD: {sheet_name}")
            for inner in range(6):
                five = np.delete(six, inner, axis=0)
                if not np.all(np.std(five, axis=0, ddof=1) > 0.0):
                    raise RuntimeError(f"RLC zero inner fit SD: {sheet_name}")
    wb.close()
    crosswalk = pd.read_csv(RLC_ROOT / "RLC1_CANONICAL_DATA.csv")
    if len(crosswalk) != 84 or crosswalk["family_key"].nunique() != 3:
        raise RuntimeError("RLC provenance crosswalk structure differs from freeze")
    if not crosswalk.groupby(["dose", "family_key"], dropna=False).size().eq(7).all():
        raise RuntimeError("RLC provenance crosswalk row support differs from freeze")
    return candidate_values, target_values

def run_rlc_family(
    candidate_values: dict[str, np.ndarray],
    target_values: dict[str, np.ndarray],
) -> tuple[pd.DataFrame, pd.DataFrame]:
    rats = [f"Rat_{i:02d}" for i in range(1, 8)]
    candidate_index = {name: i for i, name in enumerate(RLC_CANDIDATES)}
    target_family = {f"{dose}_{family}": family for dose in ("3uM", "10uM") for family in RLC_FAMILIES}
    selection_rows = []
    prediction_rows = []
    for outer_idx, outer_rat in enumerate(rats):
        outer_train = np.asarray([i for i in range(7) if i != outer_idx], dtype=int)
        inner_rats = outer_train.tolist()
        selected_by_family = {}
        selected_score_by_family = {}
        for family in RLC_FAMILIES:
            family_targets = [t for t in RLC_TARGETS if target_family[t] == family]
            score_by_candidate: dict[str, float] = {}
            error_by_candidate: dict[str, dict[int, dict[str, float]]] = {}
            for candidate in RLC_CANDIDATES:
                x = candidate_values[candidate]
                by_inner = {}
                per_rat_family = []
                for inner_idx in inner_rats:
                    fit_idx = np.asarray([i for i in outer_train if i != inner_idx], dtype=int)
                    target_errors = {}
                    for target in family_targets:
                        y = target_values[target]
                        pred_native, _, y_sd, _, _ = fit_rlc_standardized_ols(
                            x, y, fit_idx, np.asarray([inner_idx]),
                            f"{outer_rat}/{family}/{candidate}/{target}/{rats[inner_idx]}",
                        )
                        err = abs(float((y[inner_idx] - pred_native[0]) / y_sd))
                        if not math.isfinite(err):
                            raise RuntimeError("Nonfinite RLC inner score")
                        target_errors[target] = err
                    rat_mean = float(np.mean([target_errors[t] for t in family_targets]))
                    by_inner[inner_idx] = target_errors
                    per_rat_family.append(rat_mean)
                score = float(np.mean(per_rat_family))
                if not math.isfinite(score):
                    raise RuntimeError("Nonfinite RLC family candidate score")
                score_by_candidate[candidate] = score
                error_by_candidate[candidate] = by_inner
            chosen = min(RLC_CANDIDATES, key=lambda c: (score_by_candidate[c], candidate_index[c]))
            selected_by_family[family] = chosen
            selected_score_by_family[family] = score_by_candidate[chosen]
            for candidate in RLC_CANDIDATES:
                for inner_idx in inner_rats:
                    for target in family_targets:
                        selection_rows.append({
                            "outer_heldout_rat": outer_rat,
                            "family": family,
                            "candidate_action": candidate,
                            "candidate_fixed_order": candidate_index[candidate] + 1,
                            "inner_heldout_rat": rats[inner_idx],
                            "inner_fit_rats": 5,
                            "inner_target": target,
                            "inner_standardized_absolute_error": error_by_candidate[candidate][inner_idx][target],
                            "inner_rat_family_mean": float(np.mean(list(error_by_candidate[candidate][inner_idx].values()))),
                            "candidate_family_mean_score": score_by_candidate[candidate],
                            "selected_for_outer_family": candidate == chosen,
                            "tie_rule": "unrounded_binary64; exact equality uses fixed candidate order",
                        })
        for target in RLC_TARGETS:
            family = target_family[target]
            candidate = selected_by_family[family]
            x = candidate_values[candidate]
            y = target_values[target]
            pred_native, y_mu, y_sd, x_mu, x_sd = fit_rlc_standardized_ols(
                x, y, outer_train, np.asarray([outer_idx]),
                f"outer/{outer_rat}/{family}/{candidate}/{target}",
            )
            actual_z = float((y[outer_idx] - y_mu) / y_sd)
            predicted_z = float((pred_native[0] - y_mu) / y_sd)
            error_z = abs(actual_z - predicted_z)
            prediction_rows.append({
                "heldout_rat": outer_rat,
                "heldout_row_index": outer_idx + 1,
                "target": target,
                "target_dose": target.split("_", 1)[0],
                "target_family": family,
                "policy": "RLC_FAMILY",
                "candidate_action": candidate,
                "candidate_1uM_value": float(x[outer_idx]),
                "target_actual_native": float(y[outer_idx]),
                "target_predicted_native": float(pred_native[0]),
                "outer_training_target_mean": y_mu,
                "outer_training_target_sd": y_sd,
                "target_actual_standardized": actual_z,
                "target_prediction_standardized": predicted_z,
                "absolute_error_standardized": error_z,
                "absolute_error_native": abs(float(y[outer_idx] - pred_native[0])),
                "outer_training_n": 6,
                "selected_family_inner_mean_score": selected_score_by_family[family],
                "family_selector_uses_both_doses_equally": True,
            })
    selections = pd.DataFrame(selection_rows)
    predictions = pd.DataFrame(prediction_rows)
    if len(predictions) != 42 or predictions.duplicated(["heldout_rat", "target"]).any():
        raise RuntimeError("RLC family output does not contain the frozen 42 cells")
    return predictions, selections

def read_awinda_spec_frequency_grid() -> list[str]:
    spec = (ROOT / "P2_V3_ABT" / "A23_PREDECLARATION_LUNA5" / "A23_AWINDA_TARGET_SPECIFIC_RECOMPUTE_SPEC.md").read_text(encoding="utf-8")
    start = spec.index("candidate actions are exactly one matched Re/Im complex-modulus pair")
    end = spec.index("Each action costs exactly two scalar components.", start)
    values = re.findall(r"\d+\.\d+", spec[start:end])
    if len(values) != 72:
        raise RuntimeError(f"Frozen Awinda action list contains {len(values)} frequencies")
    return [f"{float(v):.4f}" for v in values]

def load_awinda_support() -> dict:
    src = V2_ROOT / "data" / "derived" / "awinda_source"
    xld = pd.read_csv(src / "WT_RLC_N47K_Control_Mavacamten_XLD_GWN_DQ1_2019_09_09.csv", low_memory=False)
    force = pd.read_csv(src / "Transgenic_mice_Fiber_Tensions_with_Fits_DQ1_2019_09_06.csv", low_memory=False)
    link = force[["AnimalID", "Fiber", "Mutation", "Treatment"]].drop_duplicates()
    if link["Fiber"].duplicated().any():
        raise RuntimeError("Awinda fibre-to-animal mapping is not one-to-one")
    files = xld[["Filename", "Mutation", "Cond"]].dropna(subset=["Filename"]).drop_duplicates()
    fmap = files.merge(link, left_on="Filename", right_on="Fiber", how="left", validate="one_to_one")
    if fmap["AnimalID"].isna().any() or not fmap["Mutation_x"].eq(fmap["Mutation_y"]).all():
        raise RuntimeError("Awinda unmapped fibre or genotype mismatch")
    fmap["drug_state"] = np.where(fmap["Cond"].eq("Control"), "Control", "0.3_uM_mavacamten")
    fmap["group"] = fmap["Mutation_x"].astype(str) + "|" + fmap["drug_state"]
    d = xld.loc[xld["pCa"].eq(4.8) & pd.to_numeric(xld["Quality"], errors="coerce").eq(1)].copy()
    d = d.merge(fmap[["Filename", "AnimalID", "Mutation_x", "drug_state", "group"]], on="Filename", how="left", validate="many_to_one")
    if d["AnimalID"].isna().any():
        raise RuntimeError("Awinda eligible spectra include an unmapped fibre")
    for c in ("ATP (mM)", "Freq (Hz)", "Em (kPa)", "Vm (kPa)"):
        d[c] = pd.to_numeric(d[c], errors="coerce")
    if d[["ATP (mM)", "Freq (Hz)", "Em (kPa)", "Vm (kPa)"]].isna().any().any():
        raise RuntimeError("Awinda eligible source rows contain nonfinite model inputs")
    freq = np.sort(d.loc[d["ATP (mM)"].eq(BASE_ATP), "Freq (Hz)"].unique())
    if len(freq) != 72 or not np.isfinite(freq).all():
        raise RuntimeError("Awinda baseline frequency count differs from the frozen 72-grid")
    freq_ids = [f"{float(f):.4f}" for f in freq]
    if freq_ids != read_awinda_spec_frequency_grid():
        raise RuntimeError("Awinda source frequency grid differs from the exact frozen action list")
    targets = AW_TARGETS
    cellcols = ["AnimalID", "Mutation_x", "drug_state", "group"]
    base_spectra = d[d["ATP (mM)"].eq(BASE_ATP)]
    base_fibres = set(base_spectra.loc[base_spectra.groupby("Filename")["Freq (Hz)"].transform("nunique").eq(72), "Filename"])
    if len(base_fibres) != 62:
        raise RuntimeError(f"Awinda baseline fibre support is {len(base_fibres)}, expected 62")
    target_fibres = {}
    for target in targets:
        td = d[d["ATP (mM)"].eq(target)]
        good = set(td.loc[td.groupby("Filename")["Freq (Hz)"].transform("nunique").eq(72), "Filename"])
        matched = base_fibres & good
        if len(matched) != 62:
            raise RuntimeError(f"Awinda target {target:g} complete fibre support differs from 62")
        target_fibres[f"{target:g}"] = len(matched)
    # Cell-level matched support: every retained cell uses identical baseline/target fibres.
    for target in targets:
        bsets = base_spectra.groupby(cellcols)["Filename"].agg(lambda x: set(x.unique()))
        tsets = d[d["ATP (mM)"].eq(target)].groupby(cellcols)["Filename"].agg(lambda x: set(x.unique()))
        if set(bsets.index) != set(tsets.index) or any(bsets[k] != tsets[k] for k in bsets.index):
            raise RuntimeError(f"Awinda baseline/target fibre membership mismatch at {target:g} mM")
    means = d.groupby(cellcols + ["ATP (mM)", "Freq (Hz)"], as_index=False).agg(
        Re=("Em (kPa)", "mean"), Im=("Vm (kPa)", "mean"), n_fibres=("Filename", "nunique")
    )
    cell_keys = []
    X = []
    Y = {target: [] for target in targets}
    per_cell_fibre_counts = {"baseline": [], **{f"{target:g}": [] for target in targets}}
    for key, sub in means.groupby(cellcols, sort=True):
        meta = dict(zip(cellcols, key))
        cell_keys.append(meta)
        baseline = sub[sub["ATP (mM)"].eq(BASE_ATP)].set_index("Freq (Hz)").reindex(freq)
        if len(baseline) != 72 or baseline[["Re", "Im"]].isna().any().any():
            raise RuntimeError(f"Awinda incomplete baseline cell: {key}")
        if baseline["n_fibres"].nunique() != 1:
            raise RuntimeError(f"Awinda baseline fibre count changes across frequency: {key}")
        per_cell_fibre_counts["baseline"].append(int(baseline["n_fibres"].iloc[0]))
        X.append(np.stack((baseline["Re"].to_numpy(float), baseline["Im"].to_numpy(float)), axis=1))
        for target in targets:
            z = sub[sub["ATP (mM)"].eq(target)].set_index("Freq (Hz)").reindex(freq)
            if len(z) != 72 or z[["Re", "Im"]].isna().any().any():
                raise RuntimeError(f"Awinda incomplete target spectrum: {key}; {target:g} mM")
            if z["n_fibres"].nunique() != 1:
                raise RuntimeError(f"Awinda target fibre count changes across frequency: {key}; {target:g} mM")
            per_cell_fibre_counts[f"{target:g}"].append(int(z["n_fibres"].iloc[0]))
            Y[target].append(np.r_[z["Re"].to_numpy(float), z["Im"].to_numpy(float)])
    cells = pd.DataFrame(cell_keys)
    Xarr = np.asarray(X, dtype=float)
    Yarr = {target: np.asarray(rows, dtype=float) for target, rows in Y.items()}
    if len(cells) != 19 or cells["AnimalID"].nunique() != 10 or Xarr.shape != (19, 72, 2):
        raise RuntimeError("Awinda cells, mice, or baseline feature dimensions differ from freeze")
    group_mouse_counts = cells.groupby("group")["AnimalID"].nunique()
    if group_mouse_counts.min() < 3:
        raise RuntimeError("Awinda group-indicator support is absent from a frozen inner training split")
    if sum(per_cell_fibre_counts["baseline"]) != 62:
        raise RuntimeError("Awinda aggregate baseline fibre count differs from frozen support")
    for target in targets:
        if sum(per_cell_fibre_counts[f"{target:g}"]) != 62:
            raise RuntimeError(f"Awinda aggregate target fibre count differs at {target:g} mM")
        if Yarr[target].shape != (19, 144) or not np.isfinite(Yarr[target]).all():
            raise RuntimeError(f"Awinda target response shape/nonfinite mismatch at {target:g} mM")
    frozen_cov = pd.read_csv(V2_ROOT / "analysis" / "03_awinda" / "continuum" / "results" / "U_TARGET_COVERAGE.csv")
    needed = ["baseline_fibres", "target_complete_fibres", "shared_complete_fibres", "animals", "baseline_cells", "target_cells", "baseline_frequencies"]
    if len(frozen_cov) != 6 or not all(frozen_cov[c].nunique() == 1 for c in needed):
        raise RuntimeError("Awinda frozen coverage reference shape differs")
    expected = {"baseline_fibres": 62, "target_complete_fibres": 62, "shared_complete_fibres": 62, "animals": 10, "baseline_cells": 19, "target_cells": 19, "baseline_frequencies": 72}
    for c, v in expected.items():
        if int(frozen_cov[c].iloc[0]) != v:
            raise RuntimeError(f"Awinda frozen coverage record mismatch: {c}")
    return {"cells": cells, "X": Xarr, "Y": Yarr, "freq": freq, "targets": targets,
            "animals": sorted(cells["AnimalID"].astype(str).unique()), "cellcols": cellcols,
            "groups": sorted(cells["group"].astype(str).unique()), "target_fibres": target_fibres,
            "per_cell_fibre_counts": per_cell_fibre_counts, "eligible_rows": len(d), "linked_files": len(fmap)}

def awinda_ridge_predict(
    data: dict,
    train_cells: np.ndarray,
    test_cells: np.ndarray,
    target: float,
    candidate_idx: int | None,
) -> np.ndarray:
    cells = data["cells"]
    group_values = cells["group"].astype(str).to_numpy()
    ordered_groups = data["groups"]
    G = np.column_stack([(group_values == g).astype(float) for g in ordered_groups])
    if candidate_idx is None:
        X_train = G[train_cells]
        X_test = G[test_cells]
        penalty = np.zeros(X_train.shape[1], dtype=float)
    else:
        features = data["X"][:, candidate_idx, :]
        mu = features[train_cells].mean(axis=0)
        sd = features[train_cells].std(axis=0, ddof=0)
        sd = np.where(sd == 0.0, 1.0, sd)
        X_train = np.column_stack((G[train_cells], (features[train_cells] - mu) / sd))
        X_test = np.column_stack((G[test_cells], (features[test_cells] - mu) / sd))
        penalty = np.r_[np.zeros(G.shape[1], dtype=float), np.full(2, ALPHA, dtype=float)]
    y_train = data["Y"][target][train_cells]
    lhs = X_train.T @ X_train + np.diag(penalty)
    beta = np.linalg.solve(lhs, X_train.T @ y_train)
    pred = X_test @ beta
    if not np.isfinite(pred).all():
        raise RuntimeError("Nonfinite Awinda ridge prediction")
    return pred

def run_awinda(data: dict) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    cells = data["cells"]
    mice = data["animals"]
    targets = data["targets"]
    freq = data["freq"]
    candidate_ids = [f"CM_ATP5_{float(f):.4f}Hz_pair" for f in freq]
    lexical_order = sorted(range(len(candidate_ids)), key=lambda i: candidate_ids[i])
    mouse_per_cell = cells["AnimalID"].astype(str).to_numpy()
    selection_rows = []
    all_cell_metrics = []
    unit_rows = []
    component_rows = []
    for outer_mouse in mice:
        outer_train_mice = [m for m in mice if m != outer_mouse]
        outer_train = np.flatnonzero(np.isin(mouse_per_cell, outer_train_mice))
        outer_test = np.flatnonzero(mouse_per_cell == outer_mouse)
        inner_mice = sorted(outer_train_mice)
        inner_scores = np.full((len(inner_mice), len(candidate_ids), len(targets)), np.nan, dtype=float)
        for inner_i, inner_mouse in enumerate(inner_mice):
            inner_train_mice = [m for m in outer_train_mice if m != inner_mouse]
            inner_train = np.flatnonzero(np.isin(mouse_per_cell, inner_train_mice))
            inner_test = np.flatnonzero(mouse_per_cell == inner_mouse)
            for candidate_i, candidate_id in enumerate(candidate_ids):
                for target_i, target in enumerate(targets):
                    pred = awinda_ridge_predict(data, inner_train, inner_test, target, candidate_i)
                    obs = data["Y"][target][inner_test]
                    denominators = np.mean(obs ** 2, axis=1)
                    if np.any(denominators <= 0.0) or not np.isfinite(denominators).all():
                        raise RuntimeError("Awinda inner spectrum has nonpositive/nonfinite RMS")
                    per_cell = np.mean((pred - obs) ** 2, axis=1) / denominators
                    inner_scores[inner_i, candidate_i, target_i] = float(np.mean(per_cell))
        if not np.isfinite(inner_scores).all():
            raise RuntimeError("Awinda inner selection score is incomplete/nonfinite")
        shared_scores = np.mean(np.mean(inner_scores, axis=2), axis=0)
        exact_scores = np.mean(inner_scores, axis=0)
        shared_i = min(lexical_order, key=lambda i: (float(shared_scores[i]), candidate_ids[i]))
        exact_is = [min(lexical_order, key=lambda i: (float(exact_scores[i, t]), candidate_ids[i])) for t in range(len(targets))]
        shared_id = candidate_ids[shared_i]
        exact_ids = [candidate_ids[i] for i in exact_is]
        for candidate_i, candidate_id in enumerate(candidate_ids):
            for inner_i, inner_mouse in enumerate(inner_mice):
                for target_i, target in enumerate(targets):
                    selection_rows.append({
                        "record_type": "INNER_MOUSE_TARGET_SCORE",
                        "outer_heldout_mouse": outer_mouse,
                        "inner_heldout_mouse": inner_mouse,
                        "policy_resolution": "SHARED_AND_EXACT_PRIMITIVE",
                        "target_ATP_mM": target,
                        "candidate_id": candidate_id,
                        "candidate_frequency_Hz": float(freq[candidate_i]),
                        "mean_cell_scaled_mse": float(inner_scores[inner_i, candidate_i, target_i]),
                        "aggregate_candidate_score": np.nan,
                        "selected": False,
                        "tie_rule": "unrounded binary64; exact equality by ascending candidate_id Unicode/code-point order",
                    })
            selection_rows.append({
                "record_type": "SHARED_CANDIDATE_SCORE",
                "outer_heldout_mouse": outer_mouse,
                "inner_heldout_mouse": "",
                "policy_resolution": "SHARED",
                "target_ATP_mM": "ALL_SIX",
                "candidate_id": candidate_id,
                "candidate_frequency_Hz": float(freq[candidate_i]),
                "mean_cell_scaled_mse": np.nan,
                "aggregate_candidate_score": float(shared_scores[candidate_i]),
                "selected": candidate_i == shared_i,
                "tie_rule": "unrounded binary64; exact equality by ascending candidate_id Unicode/code-point order",
            })
            for target_i, target in enumerate(targets):
                selection_rows.append({
                    "record_type": "EXACT_TARGET_CANDIDATE_SCORE",
                    "outer_heldout_mouse": outer_mouse,
                    "inner_heldout_mouse": "",
                    "policy_resolution": "EXACT_TARGET",
                    "target_ATP_mM": target,
                    "candidate_id": candidate_id,
                    "candidate_frequency_Hz": float(freq[candidate_i]),
                    "mean_cell_scaled_mse": np.nan,
                    "aggregate_candidate_score": float(exact_scores[candidate_i, target_i]),
                    "selected": candidate_i == exact_is[target_i],
                    "tie_rule": "unrounded binary64; exact equality by ascending candidate_id Unicode/code-point order",
                })
        selection_rows.append({
            "record_type": "SELECTED_ACTION",
            "outer_heldout_mouse": outer_mouse, "inner_heldout_mouse": "",
            "policy_resolution": "SHARED", "target_ATP_mM": "ALL_SIX", "candidate_id": shared_id,
            "candidate_frequency_Hz": float(freq[shared_i]), "mean_cell_scaled_mse": np.nan,
            "aggregate_candidate_score": float(shared_scores[shared_i]), "selected": True,
            "tie_rule": "unrounded binary64; exact equality by ascending candidate_id Unicode/code-point order",
        })
        for target_i, target in enumerate(targets):
            selection_rows.append({
                "record_type": "SELECTED_ACTION",
                "outer_heldout_mouse": outer_mouse, "inner_heldout_mouse": "",
                "policy_resolution": "EXACT_TARGET", "target_ATP_mM": target, "candidate_id": exact_ids[target_i],
                "candidate_frequency_Hz": float(freq[exact_is[target_i]]), "mean_cell_scaled_mse": np.nan,
                "aggregate_candidate_score": float(exact_scores[exact_is[target_i], target_i]), "selected": True,
                "tie_rule": "unrounded binary64; exact equality by ascending candidate_id Unicode/code-point order",
            })
        # Outer group-only baseline and the selected shared/exact fits.
        for target_i, target in enumerate(targets):
            c0_pred = awinda_ridge_predict(data, outer_train, outer_test, target, None)
            shared_pred = awinda_ridge_predict(data, outer_train, outer_test, target, shared_i)
            exact_pred = awinda_ridge_predict(data, outer_train, outer_test, target, exact_is[target_i])
            obs = data["Y"][target][outer_test]
            base_freq_ids = [f"{float(f):.4f}" for f in freq]
            for local_i, cell_i in enumerate(outer_test):
                y = obs[local_i]
                denom = float(np.mean(y ** 2))
                if not math.isfinite(denom) or denom <= 0.0:
                    raise RuntimeError("Awinda held-out spectrum has nonpositive/nonfinite RMS")
                nrmse = {
                    "C0": float(100.0 * math.sqrt(float(np.mean((c0_pred[local_i] - y) ** 2))) / math.sqrt(denom)),
                    "AW_SHARED": float(100.0 * math.sqrt(float(np.mean((shared_pred[local_i] - y) ** 2))) / math.sqrt(denom)),
                    "AW_EXACT": float(100.0 * math.sqrt(float(np.mean((exact_pred[local_i] - y) ** 2))) / math.sqrt(denom)),
                }
                if not all(math.isfinite(v) for v in nrmse.values()):
                    raise RuntimeError("Awinda outer NRMSE is nonfinite")
                meta = cells.iloc[cell_i]
                all_cell_metrics.append({
                    "heldout_mouse": outer_mouse,
                    "AnimalID": str(meta["AnimalID"]),
                    "genotype": str(meta["Mutation_x"]),
                    "drug_state": str(meta["drug_state"]),
                    "group": str(meta["group"]),
                    "target_ATP_mM": target,
                    "C0_NRMSE_pct": nrmse["C0"],
                    "shared_NRMSE_pct": nrmse["AW_SHARED"],
                    "exact_NRMSE_pct": nrmse["AW_EXACT"],
                    "shared_selected_panel_id": shared_id,
                    "exact_selected_panel_id": exact_ids[target_i],
                    "n_components": 144,
                })
                for component_i in range(144):
                    is_im = component_i >= 72
                    f_i = component_i - 72 if is_im else component_i
                    component_rows.append({
                        "heldout_mouse": outer_mouse, "AnimalID": str(meta["AnimalID"]),
                        "genotype": str(meta["Mutation_x"]), "drug_state": str(meta["drug_state"]),
                        "group": str(meta["group"]), "target_ATP_mM": target,
                        "component_index": component_i,
                        "component_type": "Im" if is_im else "Re",
                        "frequency_Hz": float(freq[f_i]), "observed": float(y[component_i]),
                        "C0_predicted": float(c0_pred[local_i, component_i]),
                        "shared_predicted": float(shared_pred[local_i, component_i]),
                        "exact_predicted": float(exact_pred[local_i, component_i]),
                        "shared_selected_panel_id": shared_id,
                        "exact_selected_panel_id": exact_ids[target_i],
                    })
    cell_metrics = pd.DataFrame(all_cell_metrics)
    unit_rows = []
    for (mouse, target), z in cell_metrics.groupby(["heldout_mouse", "target_ATP_mM"], sort=True):
        if len(z) == 0:
            raise RuntimeError("Awinda mouse/target has no stratum cells")
        if z["group"].nunique() != len(z):
            raise RuntimeError("Awinda mouse/target has duplicated genotype/drug cell")
        unit_rows.append({
            "heldout_mouse": mouse,
            "target_ATP_mM": target,
            "n_genotype_drug_cells": len(z),
            "C0_mean_spectrum_NRMSE_pct": float(z["C0_NRMSE_pct"].mean()),
            "shared_mean_spectrum_NRMSE_pct": float(z["shared_NRMSE_pct"].mean()),
            "exact_mean_spectrum_NRMSE_pct": float(z["exact_NRMSE_pct"].mean()),
            "C0_minus_shared_gain_pp": float((z["C0_NRMSE_pct"] - z["shared_NRMSE_pct"]).mean()),
            "C0_minus_exact_gain_pp": float((z["C0_NRMSE_pct"] - z["exact_NRMSE_pct"]).mean()),
            "shared_minus_exact_refinement_gain_pp": float((z["shared_NRMSE_pct"] - z["exact_NRMSE_pct"]).mean()),
            "shared_selected_panel_id": z["shared_selected_panel_id"].iloc[0],
            "exact_selected_panel_id": z["exact_selected_panel_id"].iloc[0],
        })
    unit = pd.DataFrame(unit_rows)
    if len(unit) != 60 or unit.duplicated(["heldout_mouse", "target_ATP_mM"]).any():
        raise RuntimeError("Awinda unit predictions do not preserve the 10 x 6 mouse-target vector")
    policy_columns = {
        "C0": "C0_mean_spectrum_NRMSE_pct",
        "AW_SHARED": "shared_mean_spectrum_NRMSE_pct",
        "AW_EXACT": "exact_mean_spectrum_NRMSE_pct",
    }
    summary_rows = []
    for policy, col in policy_columns.items():
        for target in targets:
            z = unit[unit["target_ATP_mM"] == target]
            summary_rows.append({
                "system_id": "AWINDA_MOUSE", "policy_id": policy, "summary_level": "target",
                "target_ATP_mM": target, "risk_NRMSE_pct": float(z[col].mean()),
                "median_mouse_target_NRMSE_pct": float(z[col].median()), "n_mice": int(z["heldout_mouse"].nunique()),
                "n_mouse_stratum_spectra": int(cell_metrics[cell_metrics["target_ATP_mM"] == target].shape[0]),
                "target_weight": 1.0 / 6.0, "unit": "percent; per-spectrum observed RMS denominator",
                "source_status": "A23_PREDECLARED_RECOMPUTATION",
            })
        target_means = [float(unit.loc[unit["target_ATP_mM"] == target, col].mean()) for target in targets]
        summary_rows.append({
            "system_id": "AWINDA_MOUSE", "policy_id": policy, "summary_level": "overall",
            "target_ATP_mM": "ALL_SIX", "risk_NRMSE_pct": float(np.mean(target_means)),
            "median_mouse_target_NRMSE_pct": float(unit[col].median()), "n_mice": 10,
            "n_mouse_stratum_spectra": int(cell_metrics.shape[0]), "target_weight": 1.0,
            "unit": "percent; equal target then equal mouse; cells averaged within mouse",
            "source_status": "A23_PREDECLARED_RECOMPUTATION",
        })
    policy_summary = pd.DataFrame(summary_rows)
    contrast_rows = []
    for r in unit.itertuples(index=False):
        for name, value in (
            ("C0_MINUS_SHARED", r.C0_minus_shared_gain_pp),
            ("C0_MINUS_EXACT", r.C0_minus_exact_gain_pp),
            ("SHARED_MINUS_EXACT", r.shared_minus_exact_refinement_gain_pp),
        ):
            contrast_rows.append({
                "contrast_id": name, "heldout_mouse": r.heldout_mouse,
                "target_ATP_mM": r.target_ATP_mM, "paired_contrast_pp": value,
                "n_genotype_drug_cells_averaged_within_mouse": r.n_genotype_drug_cells,
                "positive_value_means": "left-hand policy has higher risk; refinement gain is positive when finer policy lowers risk",
            })
    contrasts = pd.DataFrame(contrast_rows)
    selections = pd.DataFrame(selection_rows)
    components = pd.DataFrame(component_rows)
    return components, selections, policy_summary, unit, contrasts
